"""Tests for Excel output highlighting (spec 50)."""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import tempfile
import unittest
import openpyxl

from src.models import Doctor
from src.writer_excel import write_excel, CONFLICT_FILL, INFERRED_FILL, PHONETIC_FILL


class TestOutputMarkers(unittest.TestCase):
    def _write_and_read(self, doctors, warnings=None):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        write_excel(doctors, warnings or [], path)
        wb = openpyxl.load_workbook(path)
        ws = wb["Annuaire Complet"]
        os.unlink(path)
        return ws

    def test_legend_row_present(self):
        ws = self._write_and_read([Doctor(nom="DUPONT", ville="Paris")])
        self.assertEqual(ws.cell(row=1, column=1).value, "Legende:")
        # Header is on row 2, data from row 3
        self.assertEqual(ws.cell(row=2, column=1).value, "Region")

    def test_conflict_cell_highlighted(self):
        doc = Doctor(nom="DUPONT", ville="Paris", hopital="CHU A")
        doc.markers["hopital"] = {
            "kind": "conflict",
            "comment": "Conflit: voir aussi CHU B",
        }
        ws = self._write_and_read([doc])
        # Row 3 = first data row; hopital column = 11
        cell = ws.cell(row=3, column=11)
        self.assertEqual(cell.fill.start_color.rgb[-6:], CONFLICT_FILL.start_color.rgb[-6:])
        self.assertIsNotNone(cell.comment)
        self.assertIn("Conflit", cell.comment.text)

    def test_inferred_cell_highlighted(self):
        doc = Doctor(nom="DUPONT", ville="Paris", email="sec@chu.fr",
                     hopital="CHU Paris")
        doc.markers["email"] = {"kind": "inferred",
                                 "comment": "Infere depuis secretariat CHU Paris"}
        ws = self._write_and_read([doc])
        cell = ws.cell(row=3, column=14)  # email is column 14
        self.assertEqual(cell.fill.start_color.rgb[-6:], INFERRED_FILL.start_color.rgb[-6:])
        self.assertIn("Infere", cell.comment.text)

    def test_phonetic_cell_highlighted(self):
        doc = Doctor(nom="DUPONT", ville="Paris", specialite="NEPHROLOGIE",
                     hopital="CHU Paris")
        doc.markers["specialite"] = {"kind": "phonetic",
                                      "comment": "Resolu depuis: nefrologie"}
        ws = self._write_and_read([doc])
        cell = ws.cell(row=3, column=5)  # specialite is column 5
        self.assertEqual(cell.fill.start_color.rgb[-6:], PHONETIC_FILL.start_color.rgb[-6:])

    def test_unmarked_cells_have_no_marker_fill(self):
        doc = Doctor(nom="DUPONT", ville="Paris")
        ws = self._write_and_read([doc])
        cell = ws.cell(row=3, column=11)  # hopital, no marker
        # Should not be yellow/blue/gray
        rgb = (cell.fill.start_color.rgb or "")[-6:].upper()
        self.assertNotEqual(rgb, CONFLICT_FILL.start_color.rgb[-6:].upper())
        self.assertNotEqual(rgb, INFERRED_FILL.start_color.rgb[-6:].upper())
        self.assertNotEqual(rgb, PHONETIC_FILL.start_color.rgb[-6:].upper())


if __name__ == "__main__":
    unittest.main()
