"""Generate a structured, filterable Excel output file."""

from collections import Counter
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Doctor, COLUMN_HEADERS, doctor_to_row


# Styling constants
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=10)
NOTE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# Marker fills — chosen to remain readable in grayscale printing.
CONFLICT_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
INFERRED_FILL = PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid")
PHONETIC_FILL = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

_MARKER_FILLS = {
    "conflict": CONFLICT_FILL,
    "inferred": INFERRED_FILL,
    "phonetic": PHONETIC_FILL,
    "fuzzy": PHONETIC_FILL,
}

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

# Specialty display names for the summary sheet
SPECIALTIES_ORDER = [
    "NEPHROLOGIE", "PNEUMOLOGIE", "PSYCHIATRIE", "GENETIQUE",
    "OPHTALMOLOGIE", "DERMATOLOGIE", "CARDIOLOGIE", "NEUROLOGIE",
    "UROLOGIE", "AUTRE",
]

# Map Doctor attribute name -> Annuaire column index (1-based)
_FIELD_TO_COL = {
    "region": 1, "ville": 2, "type_centre": 3, "responsable_centre": 4,
    "specialite": 5, "specialite_detail": 6, "secteur": 7, "titre": 8,
    "nom": 9, "prenom": 10, "hopital": 11, "adresse": 12, "telephone": 13,
    "email": 14, "notes": 15, "source": 16, "date_import": 17,
}


def write_excel(
    doctors: list[Doctor],
    warnings: list[str],
    output_path: str,
):
    """Write the structured Excel output file."""
    wb = Workbook()

    _write_annuaire(wb, doctors)
    _write_summary(wb, doctors)
    _write_log(wb, warnings)

    wb.save(output_path)
    wb.close()


def _write_annuaire(wb: Workbook, doctors: list[Doctor]):
    """Write the main 'Annuaire Complet' sheet with highlighted markers."""
    ws = wb.active
    ws.title = "Annuaire Complet"

    # Row 1: legend explaining the marker colors
    _write_legend(ws)
    header_row = 2

    # Header row
    for col_idx, header in enumerate(COLUMN_HEADERS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sorted_doctors = sorted(doctors, key=lambda d: (d.ville, d.specialite, d.secteur, d.nom))

    for row_offset, doc in enumerate(sorted_doctors, 1):
        row_idx = header_row + row_offset
        row_data = doctor_to_row(doc)
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Per-field markers (conflict / inferred / phonetic)
        for field_name, marker in (doc.markers or {}).items():
            col = _FIELD_TO_COL.get(field_name)
            if col is None:
                continue
            cell = ws.cell(row=row_idx, column=col)
            fill = _MARKER_FILLS.get(marker.get("kind"))
            if fill:
                cell.fill = fill
            comment_text = marker.get("comment")
            if comment_text:
                cell.comment = Comment(comment_text, "ASTB")

        # Whole-row soft-highlight for notes-only records (pas de referent, idem)
        if doc.notes and any(kw in doc.notes.lower() for kw in ["pas de", "idem", "meme medecin"]):
            for col_idx in range(1, len(COLUMN_HEADERS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                # Don't overwrite a marker fill
                if cell.fill.start_color.rgb in (None, "00000000", "FFFFFFFF"):
                    cell.fill = NOTE_FILL

    # Format phone and email columns as text
    phone_col = COLUMN_HEADERS.index("Telephone") + 1
    email_col = COLUMN_HEADERS.index("Email") + 1
    for row_idx in range(header_row + 1, header_row + 1 + len(sorted_doctors)):
        ws.cell(row=row_idx, column=phone_col).number_format = "@"
        ws.cell(row=row_idx, column=email_col).number_format = "@"

    _auto_fit_columns(ws, header_row + len(sorted_doctors))

    # Freeze below the header, auto-filter across the header row
    ws.freeze_panes = f"A{header_row + 1}"
    if sorted_doctors:
        last_col = get_column_letter(len(COLUMN_HEADERS))
        ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(sorted_doctors)}"


def _write_legend(ws):
    """Write a compact one-row legend above the header explaining colors."""
    legend = [
        ("Legende:", None),
        ("Conflit hopital (revue manuelle)", CONFLICT_FILL),
        ("Valeur inferee depuis un confrere", INFERRED_FILL),
        ("Specialite resolue phonetiquement", PHONETIC_FILL),
    ]
    for col_idx, (text, fill) in enumerate(legend, 1):
        cell = ws.cell(row=1, column=col_idx, value=text)
        cell.font = Font(name="Calibri", size=9, italic=True, color="595959")
        if fill is not None:
            cell.fill = fill
        cell.alignment = Alignment(vertical="center")


def _write_summary(wb: Workbook, doctors: list[Doctor]):
    """Write the 'Par Ville' summary sheet."""
    ws = wb.create_sheet("Par Ville")

    cities = sorted(set(d.ville for d in doctors))
    specialties_present = sorted(set(d.specialite for d in doctors if d.specialite),
                                  key=lambda s: SPECIALTIES_ORDER.index(s) if s in SPECIALTIES_ORDER else 99)

    headers = ["Ville", "Total"] + specialties_present
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    counts = Counter((d.ville, d.specialite) for d in doctors if d.nom)
    city_totals = Counter(d.ville for d in doctors if d.nom)

    for row_idx, city in enumerate(cities, 2):
        ws.cell(row=row_idx, column=1, value=city).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=row_idx, column=2, value=city_totals[city]).font = BODY_FONT

        for col_idx, spec in enumerate(specialties_present, 3):
            count = counts.get((city, spec), 0)
            cell = ws.cell(row=row_idx, column=col_idx, value=count if count else "")
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="center")

    total_row = len(cities) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Calibri", size=10, bold=True)
    ws.cell(row=total_row, column=2, value=sum(city_totals.values())).font = Font(name="Calibri", size=10, bold=True)
    for col_idx, spec in enumerate(specialties_present, 3):
        total = sum(counts.get((city, spec), 0) for city in cities)
        ws.cell(row=total_row, column=col_idx, value=total).font = Font(name="Calibri", size=10, bold=True)

    _auto_fit_columns(ws, total_row)
    ws.freeze_panes = "B2"


def _write_log(wb: Workbook, warnings: list[str]):
    """Write the 'Journal Import' log sheet."""
    ws = wb.create_sheet("Journal Import")

    ws.cell(row=1, column=1, value="Message").font = HEADER_FONT
    ws.cell(row=1, column=1).fill = HEADER_FILL

    if warnings:
        for row_idx, warning in enumerate(warnings, 2):
            ws.cell(row=row_idx, column=1, value=warning).font = BODY_FONT
    else:
        ws.cell(row=2, column=1, value="Aucun avertissement").font = BODY_FONT

    ws.column_dimensions["A"].width = 100


def _auto_fit_columns(ws, max_row: int):
    """Auto-fit column widths based on content, with min/max bounds."""
    for col_idx in range(1, ws.max_column + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(1, min(max_row + 1, 50)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                lines = str(cell.value).split("\n")
                length = max(len(line) for line in lines)
                max_length = max(max_length, length)
        width = min(max(max_length + 2, 8), 40)
        ws.column_dimensions[col_letter].width = width
