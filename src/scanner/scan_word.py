"""Word file scanner: classify layout, detect fields, buffer data."""

import re
from docx import Document

from ..normalize import clean_text
from .template import (
    FileTemplate, SheetTemplate, FileFormat, LayoutType,
    FieldMapping, SectionDetector, SectionDetectorKind,
)
from .synonyms import match_field_synonym


def scan_word(path: str) -> tuple[FileTemplate, dict]:
    """Scan a Word document. Returns (template, buffer).

    buffer: {"paragraphs": [(style, text), ...], "tables": [[[cell_text]]]}
    """
    doc = Document(path)
    template = FileTemplate(file_format=FileFormat.WORD, path=path)

    paragraphs = []
    for para in doc.paragraphs:
        text = clean_text(para.text)
        style = para.style.name if para.style else ""
        paragraphs.append((style, text))

    tables_data = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([clean_text(cell.text) for cell in row.cells])
        tables_data.append(rows)

    buffer = {"paragraphs": paragraphs, "tables": tables_data}

    # Classify layout
    styled_headers = [(i, s, t) for i, (s, t) in enumerate(paragraphs)
                      if s == "List Paragraph" and t]
    heading_headers = [(i, s, t) for i, (s, t) in enumerate(paragraphs)
                       if s.startswith("Heading") and t]

    # Check tables first
    if tables_data and not styled_headers:
        sheet = _scan_tables(tables_data, template.scan_warnings)
        template.sheets.append(sheet)

    elif styled_headers:
        sheet = SheetTemplate(name="paragraphs", layout=LayoutType.PARAGRAPH_SECTIONED)
        sheet.section_detectors.append(SectionDetector(
            kind=SectionDetectorKind.STYLE_BASED,
            level="city",
            pattern="List Paragraph",
        ))
        _detect_paragraph_patterns(sheet, paragraphs)
        template.sheets.append(sheet)

    elif heading_headers:
        sheet = SheetTemplate(name="paragraphs", layout=LayoutType.PARAGRAPH_SECTIONED)
        heading_style = heading_headers[0][1]
        sheet.section_detectors.append(SectionDetector(
            kind=SectionDetectorKind.STYLE_BASED,
            level="city",
            pattern=heading_style,
        ))
        _detect_paragraph_patterns(sheet, paragraphs)
        template.sheets.append(sheet)

    else:
        sheet = SheetTemplate(name="freetext", layout=LayoutType.PARAGRAPH_FREETEXT)
        _detect_paragraph_patterns(sheet, paragraphs)
        template.sheets.append(sheet)

    return template, buffer


def _scan_tables(tables_data: list, warnings: list[str]) -> SheetTemplate:
    """Scan Word tables for column headers."""
    sheet = SheetTemplate(name="tables", layout=LayoutType.TABULAR)

    if not tables_data or not tables_data[0]:
        return sheet

    # Check first table's first row for headers
    header_row = tables_data[0][0]
    for col_idx, cell_text in enumerate(header_row):
        field_name = match_field_synonym(cell_text)
        if field_name:
            sheet.field_mappings.append(FieldMapping(
                doctor_field=field_name,
                source_label=cell_text,
                column_index=col_idx,
            ))
    sheet.data_start_row = 1
    sheet.column_count = len(header_row)

    if not sheet.field_mappings:
        warnings.append("[Scanner/Word] Tableau sans en-tetes reconnus")
        sheet.layout = LayoutType.PARAGRAPH_FREETEXT

    return sheet


def _detect_paragraph_patterns(sheet: SheetTemplate, paragraphs: list[tuple[str, str]]):
    """Detect name, phone, email patterns in paragraphs."""
    all_text = " ".join(t for _, t in paragraphs if t)

    if re.search(r"(Dr|Pr|Professeur)\s+[\w\-À-ÿ]+", all_text, re.IGNORECASE):
        sheet.field_mappings.append(FieldMapping(
            doctor_field="nom",
            source_label="Dr/Pr pattern",
            label_regex=r"(Dr|Pr|Professeur)\s+([\w\-À-ÿ]+(?:\s+[\w\-À-ÿ]+){0,3})",
        ))

    if re.search(r"(?:T[ée]l|0\d[\d\s\.]{8,})", all_text, re.IGNORECASE):
        sheet.field_mappings.append(FieldMapping(
            doctor_field="telephone",
            source_label="phone pattern",
        ))

    if re.search(r"[\w\.\-]+@[\w\.\-]+\.\w+", all_text):
        sheet.field_mappings.append(FieldMapping(
            doctor_field="email",
            source_label="email pattern",
        ))
