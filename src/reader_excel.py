"""Read a structured annuaire Excel back into Doctor objects."""

import openpyxl

from .models import Doctor, COLUMN_HEADERS


# Map column header text to Doctor field name
_HEADER_TO_FIELD = {
    "Region": "region",
    "Ville": "ville",
    "Type de centre": "type_centre",
    "Responsable du centre": "responsable_centre",
    "Specialite": "specialite",
    "Detail specialite": "specialite_detail",
    "Secteur": "secteur",
    "Titre": "titre",
    "Nom": "nom",
    "Prenom": "prenom",
    "Hopital": "hopital",
    "Adresse": "adresse",
    "Telephone": "telephone",
    "Email": "email",
    "Notes": "notes",
    "Source": "source",
    "Date import": "date_import",
}


def read_annuaire(path: str) -> tuple[list[Doctor], list[str]]:
    """Read the 'Annuaire Complet' sheet from a structured output Excel.

    Returns (doctors, warnings).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    warnings = []

    if "Annuaire Complet" not in wb.sheetnames:
        # Try the first sheet as fallback
        ws = wb[wb.sheetnames[0]]
        warnings.append(f"Feuille 'Annuaire Complet' non trouvee, utilisation de '{ws.title}'")
    else:
        ws = wb["Annuaire Complet"]

    # Find the header row (scan first few rows; writer may prepend a legend).
    col_map: dict[int, str] = {}
    header_row_idx = 1
    for candidate_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1
    ):
        mapping = {
            i: _HEADER_TO_FIELD[str(v).strip()]
            for i, v in enumerate(row)
            if v and str(v).strip() in _HEADER_TO_FIELD
        }
        if len(mapping) >= 3:  # need at least a few recognized headers
            col_map = mapping
            header_row_idx = candidate_idx
            break

    if not col_map:
        wb.close()
        return [], ["Aucun en-tete reconnu dans le fichier annuaire"]

    # Read data rows
    doctors = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        doc = Doctor()
        has_data = False
        for col_idx, field_name in col_map.items():
            if col_idx < len(row) and row[col_idx] is not None:
                value = str(row[col_idx]).strip()
                if value:
                    setattr(doc, field_name, value)
                    has_data = True

        if has_data and not doc.is_empty():
            doctors.append(doc)

    wb.close()
    return doctors, warnings
